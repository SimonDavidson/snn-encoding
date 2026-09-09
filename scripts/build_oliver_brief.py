"""Build the one-page decisions brief for Oliver as a .docx.

A companion to the encoder survey report, not a summary of it: the survey says
what has been built, this says what cannot move without Oliver. Styling helpers
are imported from the report builder rather than restated, so the two documents
cannot drift apart visually and there is one palette in the project.

Deliberately one page. Every line that is not a decision or the cost of not
taking one has been cut.

    python scripts/build_oliver_brief.py

A PDF for printing is produced alongside with

    soffice --headless --convert-to pdf --outdir reports reports/<file>.docx

which is also how the one-page claim is checked: python-docx cannot paginate,
so "one page" is verified by converting and counting /Type /Page, not asserted.

Author:        Simon Davidson & Claude
Created:       2026-09-08
Last modified: 2026-09-09
"""
import datetime as _dt
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_encoder_report import (  # noqa: E402
    AMBER, INK, LILAC, NAVY, REPORT_VERSION, ROSE, ROOT, TEAL,
    _normalise_zip_times, _rule, body, callout, caption, commit,
    open_questions, table)

#: Versioned from the second issue onward, and for the reason the overwrite
#: guard exists: the unversioned file in reports/ is the copy sent on
#: 8 September, built alongside survey v2, and it stays exactly as it was sent.
#: The suffix tracks the survey this brief accompanies, so the pair is obvious
#: in an inbox.
OUT = (ROOT / "reports" /
       f"spikeEncode_decisions_for_Oliver_v{REPORT_VERSION}.docx")


def build(force=False):
    """Same overwrite guard as the survey report, for the same reason: the
    commit hash is embedded, so a rebuild replaces a copy that may already be
    in someone's inbox."""
    if OUT.exists() and not force:
        raise SystemExit(
            f"{OUT.name} already exists. The commit hash is embedded, so "
            f"rebuilding produces a different file and would replace a copy "
            f"that may already have been sent. Run with --force to replace it "
            f"deliberately.")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)
    st.font.color.rgb = INK
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.0
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.6)
        s.top_margin = s.bottom_margin = Cm(1.3)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("spikeEncode — decisions needed from Oliver")
    r.bold = True
    r.font.size = Pt(17)
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(f"9 September 2026  ·  companion to the encoder survey, "
                  f"version {REPORT_VERSION}  ·  commit {commit()}")
    r.font.size = Pt(9.5)
    r.font.color.rgb = TEAL
    _rule(p)

    body(doc, "**Where the work is.** All six encoders implemented, known-answer suite "
              "green, and the harness running all three tasks end to end with the "
              "non-spiking reference alongside — weeks 1–2 and 5–7 of the plan, at day 21.",
              size=9.5)
    body(doc, "**What just changed.** The LDC account is accepted, so O2 — the largest "
              "blocker for three weeks — is clearing and stage one starts on TIMIT as soon "
              "as the data is in hand (D81). Until then every number is on a synthetic "
              "stand-in, whose known ground truth is what lets the pipeline controls "
              "actually fail — and they have, repeatedly. What it cannot do is evaluate "
              "control C1, the one saying the pipeline is calibrated and not merely "
              "self-consistent. **Nothing below blocks the next fortnight**: T1 and T3 run "
              "the day the data lands, and T2 only waits on a pitch tracker, since TIMIT "
              "carries no f_0 reference and §4.2 wants two of them.", size=9.5)
    callout(doc, "Survey v3 replaces every task figure in the copy you have.",
            "v2's free parameters — chiefly the label alignment — were chosen by maximising "
            "the score on the *test* split, which the protocol forbids. All five task "
            "results are re-run with them chosen inside the training split. **T1 is "
            "unaffected: the bias was exactly zero.** T2, T3 without context and P1 are "
            "not, and P1's index changes sign. Nothing about the encoders changed — only "
            "which number picked the settings. §13.4 measures it.", fill=ROSE)

    n_open = len(open_questions())
    rows = [
        ["Q43", "The corpus question you raised — TIMIT despite the accent mismatch?",
         "Answered as D81 in the reply; worth your confirmation. Nothing here is "
         "pretrained, so only a *ranking* transfers, and testing whether it survives an "
         "accent change is what stage two is for. Neither open corpus has hand-placed phone "
         "labels — T1's labels, T3's truth, and the aligner's calibration."],
        ["Q07", "Release format — are ON and OFF separate channel indices, or a polarity "
                "bit?",
         "The event file format, and controls C6 and C7 with it. Best settled before the "
         "featurisation is fixed, since doubling the channel count changes what the decoder "
         "sees. Interacts with the SHD convention, which is unipolar."],
        ["O3", "Spiketrum — who to approach, and on what terms.",
         "The primary source has been read and it changes the question (D80). The work is "
         "led from Zhejiang by Tang and colleagues, who are corresponding; Wijekoon and "
         "Alsakkal are two of eight. **Survey v2 said it was developed at Manchester by "
         "those two — that was wrong**, and it matters because the approach may need to go "
         "to Tang. The algorithm is fully published and needs no code from them, so what is "
         "at stake is attribution, not feasibility: may we implement it and report it as "
         "*our* implementation of their algorithm?"],
        ["Q31", "The mel reference does not bound T3, and v3 widens the gap.",
         "E1 scores F 0.7557 against the reference's 0.5930, and the gap *grew* under the "
         "correction above. Either the reference is mis-specified for a timing task, or "
         "§5.9 is wrong to call it an upper bound. Decides how every T3 figure is "
         "reported."],
    ]
    table(doc, ["", "Decision", "What it is holding up"], rows,
          widths=[1.2, 6.0, 10.2], size=9,
          fills={(i, 0): ROSE for i in range(len(rows))})
    caption(doc, "Not blocking today: O1 corpus specifics, needed before stage two; O4 "
                 "whether an IOP read-and-publish agreement covers the article charge; O5 "
                 "alignment strategy, which depends on O1. O2 is no longer listed: the LDC "
                 "account is accepted. Q28, whether to make the stand-in more speech-like, "
                 f"is largely settled by that — P1 cannot be answered on it either way. "
                 f"{n_open} questions are open against the design session, nine of them "
                 f"waiting on a single patch; only the rows above need you.")

    callout(doc, "One result worth a minute of the meeting.",
            "The study's central engineering question has its first number and it does not "
            "favour the answer the field assumes. The mel reference costs 128,000 bit/s; "
            "the spiking encoder reaches 98.5 per cent of its accuracy only at **398,929 "
            "bit/s, three times the representation it is approximating**, crossing near a "
            "fifth of that rate at ~91 per cent of the bound. Declared bit widths, "
            "synthetic corpus — not yet a result, but the shape of the argument the paper "
            "will have to make.", fill=LILAC)

    stamp = _dt.datetime(2026, 8, 20, 0, 0, 0)
    cp = doc.core_properties
    cp.created = cp.modified = stamp
    cp.author = cp.last_modified_by = "Simon Davidson & Claude"
    cp.title = "spikeEncode — decisions needed from Oliver"
    cp.revision = 1

    # `callout` and `table` each leave a spacer paragraph behind them, which is
    # right in the middle of a document and is a blank second page at the end of
    # a one-page one. Strip trailing empties before saving: the page count is
    # the acceptance test for this document, and it should not be failed by
    # whitespace.
    body_el = doc.element.body
    for para in reversed(doc.paragraphs):
        if para.text.strip():
            break
        body_el.remove(para._element)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    _normalise_zip_times(OUT)
    print(f"written: {OUT}")
    return OUT


if __name__ == "__main__":
    import sys
    build(force="--force" in sys.argv)
